from datetime import date

from openpyxl import load_workbook

from rechnungsscanner import excel_export
from rechnungsscanner.parser import FIELDS


def test_append_invoice_creates_file_with_header(tmp_path):
    excel_path = tmp_path / "Rechnungen.xlsx"
    data = {field: None for field in FIELDS}
    data.update(
        {
            "Dateiname": "rechnung1.pdf",
            "Unternehmensname": "Mustermann GmbH",
            "Rechnungsdatum": date(2026, 7, 1),
            "Bruttobetrag": 1190.0,
        }
    )

    excel_export.append_invoice(str(excel_path), data)

    assert excel_path.exists()
    workbook = load_workbook(excel_path)
    sheet = workbook[excel_export.SHEET_NAME]
    assert [cell.value for cell in sheet[1]] == FIELDS
    row = [cell.value for cell in sheet[2]]
    assert row[FIELDS.index("Unternehmensname")] == "Mustermann GmbH"


def test_append_invoice_appends_second_row(tmp_path):
    excel_path = tmp_path / "Rechnungen.xlsx"
    for name in ("Firma A", "Firma B"):
        data = {field: None for field in FIELDS}
        data["Unternehmensname"] = name
        excel_export.append_invoice(str(excel_path), data)

    workbook = load_workbook(excel_path)
    sheet = workbook[excel_export.SHEET_NAME]
    assert sheet.max_row == 3  # Kopfzeile + 2 Rechnungen

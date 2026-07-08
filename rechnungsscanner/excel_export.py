"""Schreibt geprüfte Rechnungsdaten als neue Zeile in eine zentrale Excel-Datei."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .parser import FIELDS

SHEET_NAME = "Rechnungen"


def _new_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(FIELDS)
    for index in range(1, len(FIELDS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 20
    return workbook


def append_invoice(excel_path: str, data: dict) -> None:
    """Fügt eine Rechnung als neue Zeile an. Legt die Datei mit Kopfzeile an, falls nötig."""
    path = Path(excel_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    else:
        workbook = _new_workbook()
        sheet = workbook[SHEET_NAME]

    row = [data.get(field) for field in FIELDS]
    sheet.append(row)

    try:
        workbook.save(path)
    except PermissionError as exc:
        raise PermissionError(
            f"Die Datei '{path.name}' ist geöffnet (z.B. in Excel) und kann nicht "
            "gespeichert werden. Bitte die Datei schließen und erneut versuchen."
        ) from exc
